"""
Render a structured TailoredResume to PDF and/or DOCX.

Rendering is deterministic and driven off the Pydantic model (not a fragile
Markdown->PDF pass), so output is consistent, ATS-friendly (real selectable
text, simple single-column layout), and has no system-level dependencies.
"""

from __future__ import annotations

import os
from typing import List

from loguru import logger
from models import JobDescription, TailoredResume


def _outfile(output_dir: str, jd: JobDescription, ext: str) -> str:
    os.makedirs(output_dir, exist_ok=True)
    return os.path.join(output_dir, f"Resume_{jd.safe_slug()}.{ext}")


# ---------------------------------------------------------------------------
# PDF (reportlab)
# ---------------------------------------------------------------------------
def render_pdf(resume: TailoredResume, jd: JobDescription, output_dir: str) -> str:
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import (
        HRFlowable,
        ListFlowable,
        ListItem,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
    )

    path = _outfile(output_dir, jd, "pdf")
    styles = getSampleStyleSheet()
    name_style = ParagraphStyle(
        "Name", parent=styles["Title"], fontSize=18, spaceAfter=2
    )
    contact_style = ParagraphStyle(
        "Contact",
        parent=styles["Normal"],
        alignment=TA_CENTER,
        fontSize=9,
        textColor="#444444",
    )
    section_style = ParagraphStyle(
        "Section",
        parent=styles["Heading2"],
        fontSize=11,
        spaceBefore=10,
        spaceAfter=2,
        textColor="#222222",
    )
    entry_head = ParagraphStyle(
        "EntryHead", parent=styles["Normal"], fontSize=10.5, spaceBefore=4, leading=13
    )
    body = ParagraphStyle("Body", parent=styles["Normal"], fontSize=9.5, leading=13)

    flow: List = [Paragraph(resume.name, name_style)]
    if resume.contact_line:
        flow.append(Paragraph(resume.contact_line, contact_style))
    flow.append(Spacer(1, 6))

    if resume.summary:
        flow.append(Paragraph("SUMMARY", section_style))
        flow.append(
            HRFlowable(width="100%", thickness=0.5, color="#cccccc", spaceAfter=4)
        )
        flow.append(Paragraph(resume.summary, body))

    for sec in resume.sections:
        flow.append(Paragraph(sec.title.upper(), section_style))
        flow.append(
            HRFlowable(width="100%", thickness=0.5, color="#cccccc", spaceAfter=4)
        )
        for entry in sec.entries:
            head_bits = [b for b in (entry.heading, entry.subheading) if b]
            head = " — ".join(
                f"<b>{b}</b>" if i == 0 else b for i, b in enumerate(head_bits)
            )
            if entry.date_range:
                head += f"  <font color='#666666'>({entry.date_range})</font>"
            if head:
                flow.append(Paragraph(head, entry_head))
            if entry.bullets:
                flow.append(
                    ListFlowable(
                        [ListItem(Paragraph(b, body)) for b in entry.bullets],
                        bulletType="bullet",
                        start="•",
                        leftIndent=12,
                    )
                )
        if sec.lines:
            for line in sec.lines:
                flow.append(Paragraph(line, body))

    SimpleDocTemplate(
        path,
        pagesize=LETTER,
        leftMargin=0.7 * inch,
        rightMargin=0.7 * inch,
        topMargin=0.6 * inch,
        bottomMargin=0.6 * inch,
    ).build(flow)
    logger.success("Wrote PDF: {}", path)
    return path


# ---------------------------------------------------------------------------
# DOCX (python-docx)
# ---------------------------------------------------------------------------
def render_docx(resume: TailoredResume, jd: JobDescription, output_dir: str) -> str:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    path = _outfile(output_dir, jd, "docx")
    doc = Document()

    name_p = doc.add_paragraph()
    name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = name_p.add_run(resume.name)
    run.bold = True
    run.font.size = Pt(18)

    if resume.contact_line:
        c = doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cr = c.add_run(resume.contact_line)
        cr.font.size = Pt(9)
        cr.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

    def add_heading(text: str) -> None:
        h = doc.add_paragraph()
        hr = h.add_run(text.upper())
        hr.bold = True
        hr.font.size = Pt(11)

    if resume.summary:
        add_heading("Summary")
        doc.add_paragraph(resume.summary)

    for sec in resume.sections:
        add_heading(sec.title)
        for entry in sec.entries:
            p = doc.add_paragraph()
            if entry.heading:
                p.add_run(entry.heading).bold = True
            if entry.subheading:
                p.add_run(f" — {entry.subheading}")
            if entry.date_range:
                dr = p.add_run(f"  ({entry.date_range})")
                dr.italic = True
            for b in entry.bullets:
                doc.add_paragraph(b, style="List Bullet")
        for line in sec.lines:
            doc.add_paragraph(line)

    doc.save(path)
    logger.success("Wrote DOCX: {}", path)
    return path


def render(
    resume: TailoredResume, jd: JobDescription, output_dir: str, formats: List[str]
) -> List[str]:
    """Render all requested formats; returns the list of written paths.

    The first path is treated as the "primary" resume to upload during
    application (PDF preferred if present).
    """
    written: List[str] = []
    if "pdf" in formats:
        written.append(render_pdf(resume, jd, output_dir))
    if "docx" in formats:
        written.append(render_docx(resume, jd, output_dir))
    if not written:
        raise ValueError("RESUME_FORMATS is empty — nothing to render.")
    return written
